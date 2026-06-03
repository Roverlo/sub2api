from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "USER_REDEEM_AND_API_GUIDE_CN.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(33, 37, 41)
MUTED = RGBColor(91, 103, 112)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D7DBE2"


def set_run_font(run, size=None, color=None, bold=None, italic=None, mono=False):
    name = "Consolas" if mono else "Calibri"
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.1

    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = BLUE
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = BLUE
    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.color.rgb = DARK_BLUE

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_header_footer(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("AI 中转站客户指南")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("请勿公开 API Key")
    set_run_font(run, size=9, color=MUTED)


def add_title(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=10, after=0)
    run = p.add_run("客户使用指南")
    set_run_font(run, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=3, after=6, line=1.05)
    run = p.add_run("我们的 AI 中转站：注册、充值/兑换与 API 接入")
    set_run_font(run, size=24, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=14, line=1.2)
    run = p.add_run(
        "面向购买或试用我们中转站服务的客户：按网站实际页面完成注册、购买/兑换、创建 API 密钥，并接入到 Codex CLI、Claude Code、Gemini CLI、OpenCode 或自己的程序。"
    )
    set_run_font(run, size=12.5, color=MUTED)


def add_para(doc, text="", after=6, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_paragraph_spacing(
        p,
        before={1: 18, 2: 14, 3: 10}.get(level, 8),
        after={1: 10, 2: 7, 3: 5}.get(level, 5),
        line=1.1,
    )
    run = p.add_run(text)
    set_run_font(
        run,
        size={1: 16, 2: 13, 3: 12}.get(level, 12),
        color=BLUE if level < 3 else DARK_BLUE,
        bold=True,
    )
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=4)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_code(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=6, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=10, color=RGBColor(30, 48, 80), mono=True)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="E2E7EE", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)

    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=2)
    run = p.add_run(title)
    set_run_font(run, size=10.5, color=DARK_BLUE, bold=True)

    p = cell.add_paragraph()
    set_paragraph_spacing(p, after=0, line=1.2)
    run = p.add_run(body)
    set_run_font(run, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)

    header_row = table.rows[0]
    header_row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        cell = header_row.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.15)
        run = p.add_run(header)
        set_run_font(run, size=10.5, color=INK, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.15)
            run = p.add_run(value)
            set_run_font(run, size=10.5, color=INK)

    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_steps(doc, steps):
    for step in steps:
        add_number(doc, step)


def add_bullets(doc, items):
    for item in items:
        add_bullet(doc, item)


def build_doc():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)
    add_title(doc)

    add_callout(
        doc,
        "这不是通用 API 教程",
        "本文按我们网站的实际页面来写：注册 /register，登录 /login，购买 /purchase，兑换 /redeem，API 密钥 /keys，用量 /usage，可用渠道 /available-channels，渠道状态 /monitor。",
    )

    add_heading(doc, "适合谁使用", 1)
    add_bullets(doc, [
        "想在国内网络环境下稳定调用大模型 API 的个人或团队。",
        "想用 OpenAI、ChatGPT-5.5 / GPT-5.5、Claude、Gemini 等模型，但不想自己处理账号、网络和接口维护的人。",
        "想把大模型接入 Cursor、Codex CLI、Claude Code、Gemini CLI、OpenCode、自己的程序或自动化脚本的人。",
        "想用余额、订阅、兑换码和 API Key 管理团队用量的人。",
    ])

    add_heading(doc, "我们中转站的优势", 1)
    advantages = [
        ("免 VPN，国内环境也能更省心调用", "客户只需要访问我们的网站和 API 地址，请求由中转站转发到上游模型服务，日常调用更省心。"),
        ("一个 API Key，接入多个模型和工具", "在「API 密钥」页面创建一个 Key 后，就可以根据后台开放的渠道调用不同模型，例如 OpenAI / ChatGPT-5.5、Claude、Gemini、Antigravity 等。"),
        ("不需要自己维护上游账号和服务器", "我们负责上游账号、调度、计费、负载均衡和接口转发，客户不用从零部署 Sub2API，也不用自己处理服务器运维。"),
        ("用量清楚，可控成本", "后台可以查看余额、兑换记录、密钥用量、模型用量和订单记录，还能给 Key 设置额度、有效期、IP 限制和速率限制。"),
        ("适合开发者和团队", "一个账号可以创建多个 API Key，按工具、项目或成员区分，方便排查和管理。"),
    ]
    for title, body in advantages:
        add_heading(doc, title, 2)
        add_para(doc, body)

    add_heading(doc, "第 1 步：打开我们的网站", 1)
    add_steps(doc, [
        "打开我们提供的中转站网址：https://你的中转站域名",
        "如果还没有账号，进入注册页：https://你的中转站域名/register",
        "如果已有账号，进入登录页：https://你的中转站域名/login",
    ])
    add_table(
        doc,
        ["页面", "地址", "用途"],
        [
            ("控制台", "/dashboard", "查看账户概览和快捷入口"),
            ("API 密钥", "/keys", "创建、复制、管理 API Key"),
            ("用量记录", "/usage", "查看请求、模型和费用明细"),
            ("可用渠道", "/available-channels", "查看当前开放的模型渠道"),
            ("渠道状态", "/monitor", "查看渠道可用性和延迟"),
            ("充值/订阅", "/purchase", "购买余额或订阅套餐"),
            ("兑换码", "/redeem", "使用兑换码兑换余额、并发或试用权限"),
            ("我的订单", "/orders", "查看充值和订阅订单"),
        ],
        [1800, 2600, 4960],
        header_fill=LIGHT_GRAY,
    )

    add_heading(doc, "第 2 步：注册账号", 1)
    add_steps(doc, [
        "进入 /register 页面。",
        "输入邮箱并设置密码。",
        "如果页面显示邀请码，请填写我们提供的邀请码。",
        "如果页面显示优惠码，可填写我们提供的优惠码，享受活动赠送额度。",
        "如果页面要求验证，请完成邮箱验证或人机验证。",
        "点击「创建账号」。注册成功后进入控制台。",
    ])

    add_heading(doc, "第 3 步：购买套餐或兑换额度", 1)
    add_heading(doc, "方式 A：在线购买", 2)
    add_steps(doc, [
        "进入 /purchase 页面。",
        "选择「充值」或「订阅」。",
        "如果是充值，输入充值金额；如果是订阅，选择适合你的模型渠道或套餐。",
        "选择支付方式，创建订单并完成支付。",
        "支付完成后，到 /orders 或控制台确认到账。",
    ])

    add_heading(doc, "方式 B：使用兑换码", 2)
    add_steps(doc, [
        "进入 /redeem 页面。",
        "查看页面顶部的当前余额和并发数。",
        "在「兑换码」输入框粘贴兑换码。",
        "点击「兑换」。",
        "页面显示「兑换成功」后，确认新余额、新并发或订阅权限。",
    ])
    add_callout(
        doc,
        "兑换码说明",
        "每个兑换码通常只能使用一次，并且区分大小写。兑换码可以增加余额、并发数或试用/订阅权限。失败时请检查空格、过期和已使用状态，并截图联系客服。",
    )

    add_heading(doc, "第 4 步：创建 API 密钥", 1)
    add_steps(doc, [
        "进入 /keys 页面并点击「创建密钥」。",
        "给密钥起一个好识别的名字，例如 codex-cli、cursor-work、server-test。",
        "选择一个可用分组。分组决定这个 Key 走哪个模型渠道和费率规则。",
        "按需设置自定义密钥、IP 限制、额度限制、速率限制和密钥有效期。",
        "保存后，在密钥列表里复制 API Key。",
    ])
    add_callout(
        doc,
        "Key 安全提醒",
        "API Key 等同于你的额度使用凭证。不要发到群里，不要贴到公开代码仓库，不要在截图中完整展示。",
    )

    add_heading(doc, "第 5 步：复制 API 地址", 1)
    add_para(doc, "在 /keys 页面顶部，网站会展示 API 地址信息。如果配置了多个自定义端点，也会在这里显示。")
    add_table(
        doc,
        ["配置项", "你应该填写"],
        [
            ("Base URL / API URL", "网站展示的 API Base URL，通常形如 https://你的中转站域名/v1"),
            ("API Key", "/keys 页面复制出来的密钥"),
            ("Model", "/available-channels 或工具提示中展示的模型名"),
        ],
        [2600, 6760],
    )
    add_code(doc, "完整接口地址示例：https://你的中转站域名/v1/chat/completions")
    add_code(doc, "Base URL 示例：https://你的中转站域名/v1")

    add_heading(doc, "第 6 步：使用「使用密钥」快速配置工具", 1)
    add_para(doc, "在 /keys 页面，每个 Key 后面都有「使用密钥」按钮。点击后，网站会根据该 Key 所属分组展示对应客户端的配置示例。")
    add_bullets(doc, [
        "OpenAI / ChatGPT-5.5 渠道：Codex CLI、Codex CLI WebSocket、OpenCode 等。",
        "Claude 渠道：Claude Code、OpenCode 等。",
        "Gemini 渠道：Gemini CLI、OpenCode 等。",
        "Antigravity 渠道：Claude Code、Gemini CLI、OpenCode 等。",
    ])

    add_heading(doc, "OpenAI 兼容接口调用示例", 1)
    add_para(doc, "请把域名、Key 和模型替换成你网站后台实际显示的内容。")
    add_heading(doc, "PowerShell 测试", 2)
    add_code(doc, '$body = @{ model = "gpt-5.5"; messages = @(@{ role = "user"; content = "你好，请用一句话介绍你自己。" }) } | ConvertTo-Json -Depth 5')
    add_code(doc, 'Invoke-RestMethod -Uri "https://你的中转站域名/v1/chat/completions" -Method Post -ContentType "application/json" -Headers @{ Authorization = "Bearer sk-你的APIKey" } -Body $body')
    add_heading(doc, "Python 示例", 2)
    add_code(doc, 'from openai import OpenAI')
    add_code(doc, 'client = OpenAI(api_key="sk-你的APIKey", base_url="https://你的中转站域名/v1")')
    add_code(doc, 'response = client.chat.completions.create(model="gpt-5.5", messages=[{"role": "user", "content": "你好，请用一句话介绍你自己。"}])')
    add_code(doc, 'print(response.choices[0].message.content)')
    add_callout(
        doc,
        "模型名以后台为准",
        "gpt-5.5 只是示例模型名。客户能使用哪些模型，请以 /available-channels 和后台模型列表为准。",
    )

    add_heading(doc, "查看用量和余额", 1)
    add_bullets(doc, [
        "/dashboard：账户概览、余额、快捷入口。",
        "/usage：请求记录、模型用量、费用明细。",
        "/keys：每个 API Key 的今日用量、总用量、额度进度和状态。",
        "/key-usage：不登录也可用 API Key 查询用量，适合给团队成员自查。",
        "/orders：查看充值或订阅订单。",
        "/redeem：查看兑换历史。",
    ])

    add_heading(doc, "常见问题", 1)
    faqs = [
        ("兑换码不能用怎么办？", "确认兑换码没有多复制空格、没有大小写错误、没有过期，也没有被使用过。如果仍然失败，请截图发给客服。"),
        ("余额到账了，为什么工具还是不能调用？", "请检查 /keys 页面是否已经创建 API Key，并确认 Key 已分配分组。没有分组的 Key 不能正确展示对应使用配置。"),
        ("为什么提示 401 或 Key 无效？", "通常是 API Key 复制错了、Key 已禁用、Key 已删除、Key 过期，或者请求头没有写成 Authorization: Bearer sk-你的APIKey。"),
        ("为什么提示模型不存在？", "模型名必须和后台开放的模型名一致。请在 /available-channels 或「使用密钥」弹窗里复制模型名，不要手打猜测。"),
        ("为什么请求超时或失败？", "可以先看 /monitor 渠道状态。如果某个渠道临时异常，可以切换到其他可用渠道，或联系客服确认当前模型状态。"),
        ("我需要 VPN 吗？", "正常使用我们的网站和 API 中转服务时，通常不需要你自己准备 VPN。你只需要能访问我们提供的网站域名和 API 地址。"),
    ]
    for question, answer in faqs:
        add_heading(doc, question, 2)
        add_para(doc, answer)

    add_heading(doc, "给客服排查时请提供", 1)
    add_steps(doc, [
        "注册邮箱或账号。",
        "当前使用的网址。",
        "兑换码失败截图，或订单号。",
        "/keys 页面密钥名称和所属分组截图。",
        "工具里填写的 Base URL 截图。",
        "报错截图。",
        "大概发生时间。",
    ])
    add_para(doc, "截图时请遮挡 API Key，只保留前后 4 到 6 位用于识别即可。")

    add_heading(doc, "安全使用提醒", 1)
    add_bullets(doc, [
        "不要公开 API Key。",
        "不要把 Key 写进公开 GitHub 仓库。",
        "不要在截图、直播、录屏里完整展示 Key。",
        "给服务器使用的 Key 建议设置 IP 白名单。",
        "给临时测试使用的 Key 建议设置有效期和额度限制。",
        "发现异常用量时，先禁用旧 Key，再联系平台客服。",
        "不要使用本服务生成、传播或处理违法违规内容。",
    ])

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
